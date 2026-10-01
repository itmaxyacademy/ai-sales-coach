import os
import http.client
import ipaddress
import re
import socket
import tempfile
from urllib.parse import urljoin, urlparse
from pathlib import Path
from typing import Optional

import fitz  # PyMuPDF
import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from markdown import markdown as md_to_html
from striprtf.striprtf import rtf_to_text


MAX_DOWNLOAD_BYTES = 10 * 1024 * 1024


def _public_addresses(url: str) -> tuple[object, list[str]]:
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("URL harus berupa HTTPS publik tanpa kredensial.")
    try:
        port = parsed.port or 443
        answers = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    except (OSError, ValueError) as exc:
        raise ValueError("Hostname URL tidak dapat di-resolve.") from exc
    addresses = list(dict.fromkeys(answer[4][0] for answer in answers))
    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
        raise ValueError("URL mengarah ke alamat jaringan non-publik.")
    return parsed, addresses


def _download_bytes(url: str) -> bytes:
    for _ in range(6):
        parsed, addresses = _public_addresses(url)
        port = parsed.port or 443
        connection = http.client.HTTPSConnection(parsed.hostname, port, timeout=15)
        pinned_ip = addresses[0]

        # Connect to the validated IP while retaining hostname-based TLS verification and Host.
        connection._create_connection = lambda address, timeout=None, source_address=None: socket.create_connection(
            (pinned_ip, address[1]), timeout, source_address
        )
        path = parsed.path or "/"
        if parsed.query:
            path += f"?{parsed.query}"
        connection.request("GET", path, headers={"User-Agent": "SalesCoachScraper/1.0"})
        response = connection.getresponse()
        if response.status in {301, 302, 303, 307, 308}:
            location = response.getheader("Location")
            connection.close()
            if not location:
                raise ValueError("Redirect tidak memiliki tujuan.")
            url = urljoin(url, location)
            continue
        if response.status >= 400:
            connection.close()
            raise ValueError("Server tujuan menolak permintaan.")
        content = response.read(MAX_DOWNLOAD_BYTES + 1)
        connection.close()
        if len(content) > MAX_DOWNLOAD_BYTES:
            raise ValueError("Ukuran dokumen melebihi batas 10 MB.")
        return content
    raise ValueError("URL terlalu banyak melakukan redirect.")


def _clean_text(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n +", "\n", text)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _extract_from_content(content: bytes, filename: str) -> str:
    ext = Path(filename.split("?", 1)[0]).suffix.lower()

    if ext in {".txt", ".md", ".csv", ".tsv", ".json", ".log"}:
        text = content.decode("utf-8", errors="ignore")
        if ext == ".md":
            text = "\n".join(BeautifulSoup(md_to_html(text), "html.parser").stripped_strings)
        return _clean_text(text)

    if ext == ".rtf":
        return _clean_text(rtf_to_text(content.decode("utf-8", errors="ignore")))

    if ext == ".html" or ext == ".htm":
        html = content.decode("utf-8", errors="ignore")
        return _clean_text("\n".join(BeautifulSoup(html, "html.parser").stripped_strings))

    if ext == ".pdf":
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            doc = fitz.open(tmp_path)
            parts = [page.get_text("text") for page in doc]
            return _clean_text("\n\n".join([p for p in parts if p.strip()]))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext == ".docx":
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            doc = Document(tmp_path)
            return _clean_text("\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext in {".csv", ".tsv"}:
        sep = "\t" if ext == ".tsv" else ","
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            df = pd.read_csv(tmp_path, sep=sep, engine="python")
            return _clean_text("\n".join(df.astype(str).agg(lambda row: " | ".join(row), axis=1).tolist()))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext in {".xlsx", ".xls", ".ods"}:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            if ext in {".xlsx", ".xls"}:
                df = pd.read_excel(tmp_path)
            else:
                df = pd.read_excel(tmp_path, engine="odf")
            return _clean_text("\n".join(df.astype(str).agg(lambda row: " | ".join(row), axis=1).tolist()))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext in {".pptx", ".ppt"}:
        try:
            import pptx
        except ImportError as exc:
            raise ValueError("python-pptx belum terpasang") from exc
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            prs = pptx.Presentation(tmp_path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
            return _clean_text("\n".join(texts))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext in {".odt"}:
        try:
            from odf import teletype
            from odf.opendocument import load
        except ImportError as exc:
            raise ValueError("odfpy belum terpasang") from exc
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            doc = load(tmp_path)
            return _clean_text(teletype.extractText(doc))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if ext in {".epub"}:
        try:
            import ebooklib
            from ebooklib import epub
        except ImportError as exc:
            raise ValueError("ebooklib belum terpasang") from exc
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(content)
            tmp_path = tmp.name
        try:
            book = epub.read_epub(tmp_path)
            items = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                items.append("\n".join(BeautifulSoup(item.get_content().decode("utf-8", errors="ignore"), "html.parser").stripped_strings))
            return _clean_text("\n\n".join(items))
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    raise ValueError(f"Format file {ext or 'tidak diketahui'} belum didukung")


def extract_text_from_url(file_url: str) -> str:
    """Extract text from common document formats via URL."""
    try:
        content = _download_bytes(file_url)
    except Exception as exc:
        raise ValueError(f"Gagal mengunduh file: {exc}") from exc

    return _extract_from_content(content, file_url)


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    """Extract text from common document formats from in-memory bytes."""
    return _extract_from_content(content, filename)
